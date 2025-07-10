import openpyxl



def find_outliers_in_columns(file_path, sheet_name, columns):
    workbook = openpyxl.load_workbook(file_path)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found in the Excel file.")

    sheet = workbook[sheet_name]

    column_indices = {}
    for col in sheet.iter_cols(1, sheet.max_column):
        if col[0].value in columns:
            column_indices[col[0].value] = col[0].column

    if len(column_indices) != len(columns):
        missing_columns = set(columns) - set(column_indices.keys())
        raise ValueError(f"Columns {missing_columns} not found in the sheet '{sheet_name}'.")

    data_rows = list(sheet.iter_rows(min_row=2, values_only=True))
    outliers = {}

    for column_name, column_index in column_indices.items():
        values = [row[column_index - 1] for row in data_rows if row[column_index - 1] is not None]
        mean = sum(values) / len(values)
        std_dev = (sum((x - mean) ** 2 for x in values) / len(values)) ** 0.5
        threshold = mean + 2 * std_dev

        outliers[column_name] = [row for row in data_rows if row[column_index - 1] is not None and row[column_index - 1] > threshold]

    return outliers

# Example usage
excel_file_path = '/Users/philliao/Downloads/Code/IB-Math-IA/2VAR_ATSBR_T2W/2var_atsbr_t2w_PUBLICTR.xlsx'

outliers = find_outliers_in_columns(excel_file_path, 'Outlier', ['PUBLICTR_percent', 'speed_acr'])
print("Outliers:")
for column, rows in outliers.items():
    print(f"Column: {column}")
    for row in rows:
        print(row)